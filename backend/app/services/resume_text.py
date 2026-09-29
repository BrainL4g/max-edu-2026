"""Извлечение текста из загруженного файла резюме.

Поддерживаются текстовые файлы (UTF-8), PDF (текстовый слой через pypdf)
и DOCX (python-docx). Сканы без текстового слоя и прочие бинарные форматы
не поддерживаются — на них поднимается ``InvalidDataError``.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from backend.app.core.exceptions import InvalidDataError

# Расширения, которые читаются как обычный текст.
TEXT_EXTENSIONS = frozenset(
    {".txt", ".md", ".text", ".rtf", ".csv", ".json", ".xml", ".html", ".htm"}
)
MAX_RESUME_BYTES = 5 * 1024 * 1024


def extract_resume_text(filename: str | None, content: bytes) -> str:
    """Текст резюме из содержимого файла.

    Raises:
        InvalidDataError: файл пустой, слишком большой или в неподдерживаемом формате.
    """
    if not content:
        raise InvalidDataError("Файл резюме пустой")
    if len(content) > MAX_RESUME_BYTES:
        raise InvalidDataError("Файл резюме слишком большой: максимум 5 МБ")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(content)
    elif suffix == ".docx":
        text = _docx_text(content)
    elif suffix in TEXT_EXTENSIONS or not suffix:
        text = content.decode("utf-8", errors="replace")
    else:
        raise InvalidDataError(f"Формат {suffix} не поддерживается: пришлите PDF, DOCX или TXT")
    text = text.strip()
    if not text:
        raise InvalidDataError(
            "В файле не нашлось текста: если это скан, пришлите текстовый вариант"
        )
    return text


def _pdf_text(content: bytes) -> str:
    """Текст PDF постранично (нужен установленный pypdf)."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - зависимость объявлена в requirements
        raise InvalidDataError("Разбор PDF недоступен: установите пакет pypdf") from exc
    try:
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise InvalidDataError("Не удалось прочитать PDF: файл повреждён или защищён") from exc


def _docx_text(content: bytes) -> str:
    """Текст DOCX: абзацы и таблицы (нужен установленный python-docx)."""
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - зависимость объявлена в requirements
        raise InvalidDataError("Разбор DOCX недоступен: установите пакет python-docx") from exc
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise InvalidDataError("Не удалось прочитать DOCX: файл повреждён") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
